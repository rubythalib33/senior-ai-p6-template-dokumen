"""Bangun 6 template .docx (docxtpl/Jinja) untuk "Pusat Skill Back-office" — Pertemuan 6.

Template dibangun dari kode supaya bisa direproduksi dan ditinjau lewat diff. Layout sepenuhnya
ditentukan di sini; LLM hanya mengisi field (function calling), program yang merender.

Jalankan dari root repo:  python sumber/bangun_template.py
"""
import os
import re
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "templates")
LOGO = os.path.join(ROOT, "sumber", "logo.png")

PERUSAHAAN = "PT Arunika Data Nusantara"
ALAMAT = "Gedung Arunika Lt. 7, Jl. Contoh Raya No. 17, Jakarta Selatan 12950"
KONTAK = "Telp (021) 555-0170  ·  sekretariat@arunika.example  ·  www.arunika.example"
CATATAN_KAKI = ("Dokumen latihan Senior AI Engineering Series. " + PERUSAHAAN +
                " adalah perusahaan fiktif; seluruh nama, nomor, dan angka bersifat ilustratif.")

FONT = "Calibri"
GELAP = RGBColor(0x1F, 0x29, 0x37)
AKSEN = RGBColor(0xB4, 0x6A, 0x0C)
ABU = RGBColor(0x6B, 0x72, 0x80)
GARIS = "9CA3AF"
LATAR_HEADER = "1F2937"
LATAR_ZEBRA = "F3F4F6"
LATAR_TOTAL = "FEF3C7"


# ---------------------------------------------------------------- utilitas XML
def _set(el, tag, **attr):
    sub = el.find(qn(tag))
    if sub is None:
        sub = OxmlElement(tag)
        el.append(sub)
    for k, v in attr.items():
        sub.set(qn(k), str(v))
    return sub


def latar_sel(cell, warna):
    tcPr = cell._tc.get_or_add_tcPr()
    _set(tcPr, "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": warna})


def margin_sel(cell, atas=60, bawah=60, kiri=90, kanan=90):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = _set(tcPr, "w:tcMar")
    for sisi, nilai in (("top", atas), ("bottom", bawah), ("start", kiri), ("end", kanan)):
        _set(mar, f"w:{sisi}", **{"w:w": nilai, "w:type": "dxa"})


def garis_tabel(tabel, warna=GARIS, ukuran=4, dalam=True):
    tblPr = tabel._tbl.tblPr
    b = _set(tblPr, "w:tblBorders")
    sisi = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if dalam else [])
    for s in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        if s in sisi:
            _set(b, f"w:{s}", **{"w:val": "single", "w:sz": ukuran, "w:space": 0, "w:color": warna})
        else:
            _set(b, f"w:{s}", **{"w:val": "nil"})


def tanpa_garis(tabel):
    tblPr = tabel._tbl.tblPr
    b = _set(tblPr, "w:tblBorders")
    for s in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        _set(b, f"w:{s}", **{"w:val": "nil"})


def garis_bawah_paragraf(p, warna="1F2937", ukuran=12):
    pPr = p._p.get_or_add_pPr()
    bdr = _set(pPr, "w:pBdr")
    _set(bdr, "w:bottom", **{"w:val": "single", "w:sz": ukuran, "w:space": 4, "w:color": warna})


def lebar_kolom(tabel, lebar_cm):
    tabel.autofit = False
    tblPr = tabel._tbl.tblPr
    _set(tblPr, "w:tblLayout", **{"w:type": "fixed"})
    grid = tabel._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(lebar_cm[i] * 567)))
    for row in tabel.rows:
        for i, cell in enumerate(row.cells):
            if i < len(lebar_cm):
                cell.width = Cm(lebar_cm[i])


def ulang_header(row):
    trPr = row._tr.get_or_add_trPr()
    _set(trPr, "w:tblHeader", **{"w:val": "true"})


def jangan_pecah(row):
    trPr = row._tr.get_or_add_trPr()
    _set(trPr, "w:cantSplit", **{"w:val": "true"})


def tinggi_baris(row, cm):
    trPr = row._tr.get_or_add_trPr()
    _set(trPr, "w:trHeight", **{"w:val": int(cm * 567), "w:hRule": "atLeast"})


def field_halaman(run, instr):
    for jenis, teks in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if jenis:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), jenis)
            run._r.append(fc)
        elif teks == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = teks
            run._r.append(t)


# ---------------------------------------------------------------- utilitas tulis
def tulis(par, teks, tebal=False, miring=False, ukuran=None, warna=None):
    r = par.add_run(teks)
    r.bold, r.italic = tebal, miring
    r.font.name = FONT
    if ukuran:
        r.font.size = Pt(ukuran)
    if warna is not None:
        r.font.color.rgb = warna
    return r


def paragraf(doc_or_cell, teks="", tebal=False, miring=False, ukuran=None, warna=None,
             rata=None, sebelum=0, sesudah=4, spasi=1.1):
    p = doc_or_cell.add_paragraph()
    if teks:
        tulis(p, teks, tebal, miring, ukuran, warna)
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(sebelum), Pt(sesudah), spasi
    if rata:
        p.alignment = rata
    return p


def isi_sel(cell, teks, tebal=False, miring=False, ukuran=10, warna=None, rata=None):
    cell.text = ""
    p = cell.paragraphs[0]
    tulis(p, teks, tebal, miring, ukuran, warna)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if rata:
        p.alignment = rata
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margin_sel(cell)
    return p


def judul_bagian(doc, teks):
    p = paragraf(doc, sebelum=10, sesudah=4)
    # huruf besar hanya untuk teks biasa; tag Jinja {{ ... }} dibiarkan apa adanya
    bagian = re.split(r"(\{\{.*?\}\})", teks)
    tulis(p, "".join(b if b.startswith("{{") else b.upper() for b in bagian), tebal=True, ukuran=10.5, warna=AKSEN)
    p.paragraph_format.keep_with_next = True
    return p


def baris_tag(tabel, tag):
    """Baris berisi tag docxtpl ({%tr ... %}); baris ini dihapus saat render."""
    row = tabel.add_row()
    row.cells[0].text = tag
    return row


# ---------------------------------------------------------------- kerangka dokumen
def dokumen_baru():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(2.2), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)

    # kop surat
    kop = doc.add_table(rows=1, cols=2)
    tanpa_garis(kop)
    lebar_kolom(kop, [2.0, 14.8])
    c0, c1 = kop.rows[0].cells
    c0.paragraphs[0].add_run().add_picture(LOGO, width=Cm(1.6))
    c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = c1.paragraphs[0]
    tulis(p, PERUSAHAAN.upper(), tebal=True, ukuran=15, warna=GELAP)
    p.paragraph_format.space_after = Pt(0)
    p2 = c1.add_paragraph()
    tulis(p2, ALAMAT, ukuran=8.5, warna=ABU)
    p2.paragraph_format.space_after = Pt(0)
    p3 = c1.add_paragraph()
    tulis(p3, KONTAK, ukuran=8.5, warna=ABU)
    p3.paragraph_format.space_after = Pt(0)
    c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    garis = paragraf(doc, sesudah=8)
    garis_bawah_paragraf(garis, warna="B46A0C", ukuran=18)
    garis.paragraph_format.line_spacing = 0.6

    # catatan kaki + nomor halaman
    fp = sec.footer.paragraphs[0]
    tulis(fp, CATATAN_KAKI, miring=True, ukuran=7.5, warna=ABU)
    fp2 = sec.footer.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tulis(fp2, "Halaman ", ukuran=8, warna=ABU)
    field_halaman(tulis(fp2, "", ukuran=8, warna=ABU), "PAGE")
    tulis(fp2, " dari ", ukuran=8, warna=ABU)
    field_halaman(tulis(fp2, "", ukuran=8, warna=ABU), "NUMPAGES")
    return doc


def judul_dokumen(doc, judul, sub="Nomor: {{ nomor }}"):
    paragraf(doc, judul, tebal=True, ukuran=13.5, warna=GELAP, rata=WD_ALIGN_PARAGRAPH.CENTER, sesudah=0)
    paragraf(doc, sub, ukuran=10, warna=ABU, rata=WD_ALIGN_PARAGRAPH.CENTER, sesudah=10)


def tabel_info(doc, pasangan, lebar=(4.4, 12.4)):
    t = doc.add_table(rows=0, cols=3)
    tanpa_garis(t)
    for kunci, nilai in pasangan:
        row = t.add_row()
        isi_sel(row.cells[0], kunci, tebal=True, ukuran=10, warna=GELAP)
        isi_sel(row.cells[1], ":", ukuran=10)
        isi_sel(row.cells[2], nilai, ukuran=10)
        for c in row.cells:
            margin_sel(c, atas=25, bawah=25, kiri=40, kanan=40)
    lebar_kolom(t, [lebar[0], 0.4, lebar[1]])
    return t


def tabel_data(doc, kepala, sel_isi, lebar, tag_loop, rata_kanan=(), baris_total=None, ukuran=9.5):
    """Tabel bergaris dengan header gelap, baris loop docxtpl, dan baris total opsional.

    baris_total: daftar (label, ekspresi) yang ditaruh di kolom terakhir; label digabung di kolom lain.
    """
    t = doc.add_table(rows=1, cols=len(kepala))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    garis_tabel(t)
    for i, k in enumerate(kepala):
        c = t.rows[0].cells[i]
        isi_sel(c, k, tebal=True, ukuran=9, warna=RGBColor(0xFF, 0xFF, 0xFF),
                rata=WD_ALIGN_PARAGRAPH.RIGHT if i in rata_kanan else None)
        latar_sel(c, LATAR_HEADER)
    ulang_header(t.rows[0])
    baris_tag(t, "{%tr for " + tag_loop + " %}")
    row = t.add_row()
    jangan_pecah(row)
    for i, ekspr in enumerate(sel_isi):
        isi_sel(row.cells[i], ekspr, ukuran=ukuran, rata=WD_ALIGN_PARAGRAPH.RIGHT if i in rata_kanan else None)
    baris_tag(t, "{%tr endfor %}")
    for label, ekspr in (baris_total or []):
        r = t.add_row()
        jangan_pecah(r)
        gab = r.cells[0].merge(r.cells[len(kepala) - 2])
        isi_sel(gab, label, tebal=True, ukuran=9.5, rata=WD_ALIGN_PARAGRAPH.RIGHT)
        isi_sel(r.cells[len(kepala) - 1], ekspr, tebal=True, ukuran=9.5, rata=WD_ALIGN_PARAGRAPH.RIGHT)
        for c in r.cells:
            latar_sel(c, LATAR_TOTAL)
    lebar_kolom(t, lebar)
    return t


def daftar(doc, ekspr_list, isi="{{ x }}", penanda="•", var="x"):
    paragraf(doc, "{%p for " + var + " in " + ekspr_list + " %}", sesudah=0)
    p = paragraf(doc, sesudah=2)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    tulis(p, penanda + "  ", warna=AKSEN, tebal=True)
    tulis(p, isi)
    paragraf(doc, "{%p endfor %}", sesudah=0)


def daftar_bernomor(doc, ekspr_list, isi="{{ x }}", var="x"):
    daftar(doc, ekspr_list, isi=isi, penanda="{{ loop.index }}.", var=var)


def kotak_peringatan(doc, ekspr_list, judul="Perhatian"):
    paragraf(doc, "{%p if " + ekspr_list + " %}", sesudah=0)
    paragraf(doc, sesudah=0, spasi=0.5)  # jarak dari elemen di atasnya (ikut hilang bila tidak ada peringatan)
    t = doc.add_table(rows=1, cols=1)
    garis_tabel(t, warna="B46A0C", ukuran=8, dalam=False)
    c = t.rows[0].cells[0]
    latar_sel(c, LATAR_TOTAL)
    isi_sel(c, judul, tebal=True, ukuran=9.5, warna=AKSEN)
    c.add_paragraph("{%p for w in " + ekspr_list + " %}")
    p = c.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    tulis(p, "•  ", tebal=True, warna=AKSEN, ukuran=9.5)
    tulis(p, "{{ w }}", ukuran=9.5)
    c.add_paragraph("{%p endfor %}")
    lebar_kolom(t, [16.8])
    paragraf(doc, "{%p endif %}", sesudah=0)


def blok_ttd(doc, kolom, tempat_tanggal="Jakarta, {{ tanggal }}"):
    """kolom: daftar (peran, jabatan_ekspr, nama_ekspr[, bawah_ekspr]). Satu atau dua kolom tanda tangan.

    Baris: tempat/tanggal · peran · jabatan · ruang tanda tangan · nama (bergaris bawah) · teks di bawah nama.
    Seluruh blok dijaga tetap satu halaman (keep_with_next + baris tidak terpotong).
    """
    sp = paragraf(doc, sesudah=4)
    sp.paragraph_format.keep_with_next = True
    n = len(kolom)
    t = doc.add_table(rows=6, cols=n)
    tanpa_garis(t)
    if n == 1:
        lebar_kolom(t, [16.8])
    else:
        lebar_kolom(t, [16.8 / n] * n)
    for i, k in enumerate(kolom):
        peran, jab, nama = k[:3]
        bawah = k[3] if len(k) > 3 else ""
        rata = WD_ALIGN_PARAGRAPH.CENTER
        isi_sel(t.rows[0].cells[i], tempat_tanggal if i == n - 1 else "", ukuran=10, rata=rata)
        isi_sel(t.rows[1].cells[i], peran, ukuran=10, rata=rata)
        isi_sel(t.rows[2].cells[i], jab, tebal=True, ukuran=10, rata=rata)
        isi_sel(t.rows[3].cells[i], "", ukuran=10)
        p = isi_sel(t.rows[4].cells[i], nama, tebal=True, ukuran=10, rata=rata)
        p.runs[0].underline = True
        isi_sel(t.rows[5].cells[i], bawah, ukuran=10, rata=rata)
    tinggi_baris(t.rows[3], 1.8)
    for j, r in enumerate(t.rows):
        jangan_pecah(r)
        if j < len(t.rows) - 1:
            for c in r.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    if n == 1:  # satu penanda tangan: blok berada di sisi kanan halaman
        lebar_kolom(t, [16.8])
        tblPr = t._tbl.tblPr
        _set(tblPr, "w:tblInd", **{"w:w": int(9.8 * 567), "w:type": "dxa"})
        lebar_kolom(t, [7.0])
    return t


def simpan(doc, nama):
    os.makedirs(OUT, exist_ok=True)
    path = os.path.join(OUT, nama + ".docx")
    doc.save(path)
    print("ditulis", os.path.relpath(path, ROOT))


# ---------------------------------------------------------------- 1. memo pengadaan
def memo_pengadaan():
    doc = dokumen_baru()
    judul_dokumen(doc, "MEMO PERMINTAAN PENGADAAN BARANG/JASA")
    tabel_info(doc, [
        ("Kepada", "{{ kepada }}"),
        ("Dari", "{{ jabatan_pemohon }} – {{ unit_pemohon }}"),
        ("Tanggal", "{{ tanggal }}"),
        ("Perihal", "{{ perihal }}"),
        ("Jenis pengadaan", "{{ jenis_pengadaan }}"),
        ("Dibutuhkan paling lambat", "{{ kebutuhan_tanggal }}"),
        ("Sumber anggaran", "{{ sumber_anggaran }}"),
    ])
    judul_bagian(doc, "1. Latar belakang dan justifikasi")
    paragraf(doc, "{{ latar_belakang }}", rata=WD_ALIGN_PARAGRAPH.JUSTIFY)
    judul_bagian(doc, "2. Rincian kebutuhan")
    tabel_data(doc,
               ["No", "Uraian", "Spesifikasi", "Jml", "Satuan", "Harga satuan (Rp)", "Jumlah (Rp)"],
               ["{{ it.no }}", "{{ it.uraian }}", "{{ it.spesifikasi }}", "{{ it.jumlah }}", "{{ it.satuan }}",
                "{{ it.harga_satuan }}", "{{ it.jumlah_harga }}"],
               [0.8, 3.4, 4.6, 1.0, 1.5, 2.7, 2.8], "it in items", rata_kanan=(0, 3, 5, 6),
               baris_total=[("Subtotal", "{{ subtotal }}"), ("PPN {{ ppn_persen }}%", "{{ ppn }}"),
                            ("Total", "{{ total }}")])
    paragraf(doc, "Terbilang: {{ total_terbilang }}", miring=True, ukuran=9.5, sebelum=4)
    judul_bagian(doc, "3. Metode pengadaan")
    t = tabel_info(doc, [("Metode", "{{ metode_pengadaan }}"),
                         ("Penawaran pembanding", "minimal {{ min_penawaran }} penawaran tertulis"),
                         ("Dasar aturan", "{{ dasar_metode }}")])
    kotak_peringatan(doc, "peringatan")
    judul_bagian(doc, "4. Lampiran wajib")
    daftar(doc, "lampiran", isi="{{ x }}", penanda="☐")
    judul_bagian(doc, "5. Persetujuan")
    tb = tabel_data(doc, ["Tingkat", "Jabatan", "Nama", "Tanda tangan", "Tanggal"],
                    ["{{ s.tingkat }}", "{{ s.jabatan }}", "", "", ""],
                    [1.6, 5.2, 4.0, 3.4, 2.6], "s in persetujuan")
    tinggi_baris(tb.rows[2], 1.1)
    blok_ttd(doc, [("Pemohon,", "{{ jabatan_pemohon }}", "{{ nama_pemohon }}")])
    simpan(doc, "memo_pengadaan")


# ---------------------------------------------------------------- 2. SPPD
def sppd():
    doc = dokumen_baru()
    judul_dokumen(doc, "SURAT PERINTAH PERJALANAN DINAS")
    tabel_info(doc, [
        ("Pemberi tugas", "{{ pemberi_tugas }}"),
        ("Nama pegawai", "{{ nama_pegawai }}"),
        ("Jabatan / tingkat", "{{ jabatan }} / {{ tingkat }}"),
        ("Unit kerja", "{{ unit }}"),
        ("Maksud perjalanan", "{{ maksud_perjalanan }}"),
        ("Kota asal – tujuan", "{{ kota_asal }} – {{ kota_tujuan }}"),
        ("Moda transportasi", "{{ moda_transportasi }}"),
        ("Tanggal berangkat", "{{ tanggal_berangkat }}"),
        ("Tanggal kembali", "{{ tanggal_kembali }}"),
        ("Lama perjalanan", "{{ lama_hari }} hari ({{ lama_malam }} malam)"),
    ])
    judul_bagian(doc, "Rincian biaya perjalanan")
    tabel_data(doc, ["No", "Komponen", "Volume", "Satuan", "Tarif (Rp)", "Jumlah (Rp)"],
               ["{{ b.no }}", "{{ b.komponen }}", "{{ b.volume }}", "{{ b.satuan }}", "{{ b.tarif }}", "{{ b.jumlah }}"],
               [0.8, 6.0, 1.6, 2.2, 3.0, 3.2], "b in biaya", rata_kanan=(0, 2, 4, 5),
               baris_total=[("Total biaya", "{{ total }}")])
    paragraf(doc, "Terbilang: {{ total_terbilang }}", miring=True, ukuran=9.5, sebelum=4)
    kotak_peringatan(doc, "peringatan")
    judul_bagian(doc, "Ketentuan pertanggungjawaban")
    daftar_bernomor(doc, "ketentuan")
    blok_ttd(doc, [("Pegawai yang ditugaskan,", "{{ jabatan }}", "{{ nama_pegawai }}"),
                   ("Pemberi tugas,", "{{ pemberi_tugas }}", "{{ nama_pemberi_tugas }}")])
    simpan(doc, "sppd")


# ---------------------------------------------------------------- 3. nota dinas
def nota_dinas():
    doc = dokumen_baru()
    judul_dokumen(doc, "NOTA DINAS", sub="")
    tabel_info(doc, [
        ("Nomor", "{{ nomor }}"),
        ("Kepada", "{{ kepada }}"),
        ("Dari", "{{ dari }}"),
        ("Tembusan", "{{ tembusan_teks }}"),
        ("Tanggal", "{{ tanggal }}"),
        ("Sifat", "{{ sifat }}"),
        ("Lampiran", "{{ lampiran }}"),
        ("Perihal", "{{ perihal }}"),
    ], lebar=(3.0, 13.8))
    g = paragraf(doc, sesudah=8)
    garis_bawah_paragraf(g, warna=GARIS, ukuran=6)
    paragraf(doc, "{%p for x in isi %}", sesudah=0)
    paragraf(doc, "{{ x }}", rata=WD_ALIGN_PARAGRAPH.JUSTIFY, sesudah=8, spasi=1.25)
    paragraf(doc, "{%p endfor %}", sesudah=0)
    paragraf(doc, "Demikian disampaikan. Atas perhatian dan kerja samanya, kami ucapkan terima kasih.",
             rata=WD_ALIGN_PARAGRAPH.JUSTIFY, spasi=1.25)
    blok_ttd(doc, [("", "{{ jabatan_penandatangan }}", "{{ nama_penandatangan }}")])
    simpan(doc, "nota_dinas")


# ---------------------------------------------------------------- 4. notulen rapat
def notulen_rapat():
    doc = dokumen_baru()
    judul_dokumen(doc, "NOTULEN RAPAT")
    tabel_info(doc, [
        ("Judul rapat", "{{ judul_rapat }}"),
        ("Hari, tanggal", "{{ tanggal_rapat }}"),
        ("Waktu", "{{ waktu_mulai }} – {{ waktu_selesai }} WIB"),
        ("Tempat", "{{ tempat }}"),
        ("Pimpinan rapat", "{{ pimpinan_rapat }}"),
        ("Notulis", "{{ notulis }}"),
    ])
    judul_bagian(doc, "Peserta ({{ peserta|length }} orang)")
    daftar_bernomor(doc, "peserta")
    judul_bagian(doc, "Agenda")
    daftar_bernomor(doc, "agenda")
    judul_bagian(doc, "Pembahasan")
    tabel_data(doc, ["No", "Agenda", "Ringkasan pembahasan"],
               ["{{ p.no }}", "{{ p.agenda }}", "{{ p.ringkasan }}"],
               [0.8, 4.4, 11.6], "p in pembahasan", rata_kanan=(0,))
    judul_bagian(doc, "Keputusan")
    daftar_bernomor(doc, "keputusan")
    judul_bagian(doc, "Tindak lanjut")
    tabel_data(doc, ["No", "Tugas", "PIC", "Tenggat", "Catatan"],
               ["{{ t.no }}", "{{ t.tugas }}", "{{ t.pic }}", "{{ t.tenggat }}", "{{ t.catatan }}"],
               [0.8, 6.4, 3.0, 3.0, 3.6], "t in tindak_lanjut", rata_kanan=(0,))
    kotak_peringatan(doc, "peringatan")
    blok_ttd(doc, [("Notulis,", "", "{{ notulis }}"), ("Pimpinan rapat,", "", "{{ pimpinan_rapat }}")],
             tempat_tanggal="Jakarta, {{ tanggal }}")
    simpan(doc, "notulen_rapat")


# ---------------------------------------------------------------- 5. laporan insiden
def laporan_insiden():
    doc = dokumen_baru()
    judul_dokumen(doc, "LAPORAN INSIDEN LAYANAN TI")
    # pita severity: warna sel diatur program lewat tag cellbg docxtpl
    t = doc.add_table(rows=1, cols=2)
    garis_tabel(t, dalam=False)
    c0, c1 = t.rows[0].cells
    isi_sel(c0, "{% cellbg sev_warna %}{{ severity }}", tebal=True, ukuran=16,
            warna=RGBColor(0xFF, 0xFF, 0xFF), rata=WD_ALIGN_PARAGRAPH.CENTER)
    p = isi_sel(c1, "{{ severity_label }}", tebal=True, ukuran=11)
    q = c1.add_paragraph()
    tulis(q, "Dampak {{ dampak }} × urgensi {{ urgensi }}  ·  status: {{ status }}", ukuran=9.5, warna=ABU)
    lebar_kolom(t, [2.6, 14.2])
    paragraf(doc, sesudah=6)
    tabel_info(doc, [
        ("Judul insiden", "{{ judul }}"),
        ("Layanan terdampak", "{{ layanan_terdampak }}"),
        ("Mulai terjadi", "{{ waktu_mulai }}"),
        ("Terdeteksi", "{{ waktu_terdeteksi }}"),
        ("Pulih", "{{ waktu_pulih }}"),
        ("Durasi gangguan", "{{ durasi }}"),
        ("Pelapor", "{{ pelapor }}"),
    ])
    judul_bagian(doc, "Ringkasan")
    paragraf(doc, "{{ ringkasan }}", rata=WD_ALIGN_PARAGRAPH.JUSTIFY)
    judul_bagian(doc, "Dampak terhadap pengguna")
    paragraf(doc, "{{ dampak_uraian }}", rata=WD_ALIGN_PARAGRAPH.JUSTIFY)
    judul_bagian(doc, "Linimasa")
    tabel_data(doc, ["Waktu", "Kejadian"], ["{{ l.waktu }}", "{{ l.kejadian }}"], [3.4, 13.4], "l in linimasa")
    judul_bagian(doc, "Akar masalah")
    paragraf(doc, "{{ akar_masalah }}", rata=WD_ALIGN_PARAGRAPH.JUSTIFY)
    judul_bagian(doc, "Tindakan pemulihan")
    paragraf(doc, "{{ tindakan_pemulihan }}", rata=WD_ALIGN_PARAGRAPH.JUSTIFY)
    judul_bagian(doc, "Tindak lanjut pencegahan")
    tabel_data(doc, ["No", "Tugas", "PIC", "Tenggat"],
               ["{{ t.no }}", "{{ t.tugas }}", "{{ t.pic }}", "{{ t.tenggat }}"],
               [0.8, 9.4, 3.4, 3.2], "t in tindak_lanjut", rata_kanan=(0,))
    kotak_peringatan(doc, "peringatan")
    blok_ttd(doc, [("Pelapor,", "", "{{ pelapor }}"), ("Mengetahui,", "{{ jabatan_penanggung_jawab }}", "")])
    simpan(doc, "laporan_insiden")


# ---------------------------------------------------------------- 6. surat penawaran
def surat_penawaran():
    doc = dokumen_baru()
    t = doc.add_table(rows=0, cols=2)
    tanpa_garis(t)
    for kiri, kanan in (("Nomor   : {{ nomor }}", "Jakarta, {{ tanggal }}"),
                        ("Lampiran: {{ lampiran }}", ""),
                        ("Perihal  : {{ perihal }}", "")):
        r = t.add_row()
        isi_sel(r.cells[0], kiri, ukuran=10)
        isi_sel(r.cells[1], kanan, ukuran=10, rata=WD_ALIGN_PARAGRAPH.RIGHT)
    lebar_kolom(t, [11.0, 5.8])
    paragraf(doc, sesudah=6)
    paragraf(doc, "Kepada Yth.", sesudah=0)
    paragraf(doc, "{{ nama_klien }}", tebal=True, sesudah=0)
    paragraf(doc, "{{ alamat_klien }}", sesudah=0)
    paragraf(doc, "U.p. {{ up }}", sesudah=10)
    paragraf(doc, "Dengan hormat,", sesudah=6)
    paragraf(doc, "{{ pembuka }}", rata=WD_ALIGN_PARAGRAPH.JUSTIFY, spasi=1.25)
    judul_bagian(doc, "Ruang lingkup")
    daftar_bernomor(doc, "ruang_lingkup")
    judul_bagian(doc, "Rincian harga")
    tabel_data(doc, ["No", "Uraian", "Jml", "Satuan", "Harga satuan (Rp)", "Jumlah (Rp)"],
               ["{{ it.no }}", "{{ it.uraian }}", "{{ it.jumlah }}", "{{ it.satuan }}", "{{ it.harga_satuan }}",
                "{{ it.jumlah_harga }}"],
               [0.8, 6.6, 1.0, 1.8, 3.2, 3.4], "it in items", rata_kanan=(0, 2, 4, 5),
               baris_total=[("Subtotal", "{{ subtotal }}"), ("PPN {{ ppn_persen }}%", "{{ ppn }}"),
                            ("Total", "{{ total }}")])
    paragraf(doc, "Terbilang: {{ total_terbilang }}", miring=True, ukuran=9.5, sebelum=4)
    judul_bagian(doc, "Ketentuan")
    daftar_bernomor(doc, "ketentuan")
    paragraf(doc, "Demikian penawaran ini kami sampaikan. Kami berharap dapat bekerja sama dengan "
                  "{{ nama_klien }}. Atas perhatian Bapak/Ibu, kami ucapkan terima kasih.",
             rata=WD_ALIGN_PARAGRAPH.JUSTIFY, sebelum=8, spasi=1.25)
    blok_ttd(doc, [(PERUSAHAAN, "", "{{ nama_penandatangan }}", "{{ jabatan_penandatangan }}")],
             tempat_tanggal="Hormat kami,")
    simpan(doc, "surat_penawaran")


if __name__ == "__main__":
    memo_pengadaan()
    sppd()
    nota_dinas()
    notulen_rapat()
    laporan_insiden()
    surat_penawaran()
